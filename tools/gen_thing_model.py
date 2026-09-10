#!/usr/bin/env python3
"""按「兰溪水亭」模板生成 SL651 协议物模型 Excel。

工作簿结构 (与 兰溪水亭.xlsx 一致):
  - sheet 物模型:  全部设备类型及其点位字典
      列: 设备类型 | 字段名 | 中文名 | 属性类型 | 单位 | 数据类型 | 数值说明
  - sheet <设施类型> (每设施类型一页): 该设施类型包含的 设备类型/具体设备/点位
      列: 设备类型 | 具体设备 | 属性 | 中文名 | 寄存器地址
      SL651 无寄存器概念, 寄存器地址列留空
  - sheet 设施清单: 各设施及其设施类型
      列: 设施类型 | 设施名称

点位目录: bridge/src/data/points.tsv (标识符=规约源码, 全部不翻译)。
默认生成一个设施类型与名称均为「薄天雅安」的设施, 含协议支持的全部设备类型
(雨量计/水位计/流量计/.../遥测终端), 每类型一台具体设备。

用法:
  python3 tools/gen_thing_model.py [--place 薄天雅安] [--out configs/sl651-thing-model.xlsx]

依赖 openpyxl (可在 onboard-tool 环境运行: uv run --project <onboard-tool> python 本脚本)
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

REPO = Path(__file__).resolve().parent.parent
POINTS_TSV = REPO / "bridge/src/data/points.tsv"

# 整数型点位 (规约 N(D,0)); 其余数值型默认小数
INT_POINTS = {"FL", "GN", "GS", "GT", "NS", "UC", "UE", "COND", "TURB", "7A"}
STRING_POINTS = {"RGZS", "PIC_file"}

REMARKS = {
    "ZT": "SL651 ZT 状态字原始值 (表58 位图)",
    "7A": "厂商自定义 信号强度",
    "FFA0": "厂商自定义 设备温度",
    "DRP": "小时报 12 组 5 分钟时段雨量",
    "RGZS": "人工置数报 (35H) 原编码全文",
    "PIC_size": "图片报 (36H) 落盘文件大小",
    "PIC_file": "图片报 (36H) 落盘文件名",
}


def load_points() -> list[dict]:
    rows = []
    for line in POINTS_TSV.read_text(encoding="utf-8").splitlines():
        if line.startswith("#") or not line.strip():
            continue
        parts = line.split("\t")
        rows.append(
            {
                "group": parts[0],   # 设备类型
                "ident": parts[1],   # 规约源码标识符
                "zh": parts[2],
                "kind": parts[3],   # t=遥测 a=属性
                "unit": parts[4] if len(parts) > 4 else "",
            }
        )
    return rows


def dtype_of(ident: str) -> str:
    if ident in STRING_POINTS:
        return "字符串"
    if ident.startswith("ZT_"):
        return "布尔"
    if ident in INT_POINTS:
        return "整数"
    return "小数"


def main() -> int:
    ap = argparse.ArgumentParser(description="按兰溪水亭模板生成 SL651 物模型 Excel")
    ap.add_argument("--place", default="薄天雅安", help="设施类型与设施名称 (默认 薄天雅安)")
    ap.add_argument("--device-suffix", default="", help="具体设备名后缀 (默认与设备类型同名)")
    ap.add_argument("--out", default="configs/sl651-thing-model.xlsx", help="输出 xlsx 路径")
    args = ap.parse_args()

    points = load_points()
    groups: dict[str, list[dict]] = {}
    for p in points:
        groups.setdefault(p["group"], []).append(p)
    if not groups:
        ap.error(f"点位目录为空: {POINTS_TSV}")

    wb = Workbook()
    bold = Font(bold=True)

    # ---- sheet 1: 物模型 (全部设备类型 + 点位字典) ----
    ws = wb.active
    ws.title = "物模型"
    ws.append(["设备类型", "字段名", "中文名", "属性类型", "单位", "数据类型", "数值说明"])
    for c in ws[1]:
        c.font = bold
    for group, pts in groups.items():
        for p in pts:
            ws.append(
                [
                    group,
                    p["ident"],
                    p["zh"],
                    "遥测数据" if p["kind"] == "t" else "只读属性",
                    p["unit"],
                    dtype_of(p["ident"]),
                    REMARKS.get(p["ident"], ""),
                ]
            )
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["C"].width = 20
    ws.column_dimensions["G"].width = 36

    # ---- sheet 2: 设施类型页 (设备类型/具体设备/点位; 寄存器地址留空) ----
    ws2 = wb.create_sheet(args.place)
    ws2.append(["设备类型", "具体设备", "属性", "中文名", "寄存器地址"])
    for c in ws2[1]:
        c.font = bold
    for group, pts in groups.items():
        device = group + args.device_suffix  # 每设备类型一台具体设备, 与类型同名
        for p in pts:
            ws2.append([group, device, p["ident"], p["zh"], ""])
    ws2.column_dimensions["A"].width = 14
    ws2.column_dimensions["B"].width = 18
    ws2.column_dimensions["D"].width = 20

    # ---- sheet 3: 设施清单 ----
    ws3 = wb.create_sheet("设施清单")
    ws3.append(["设施类型", "设施名称"])
    for c in ws3[1]:
        c.font = bold
    ws3.append([args.place, args.place])

    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)

    print(f"已生成 {out}")
    print(f"  设施: {args.place} (设施类型同名), 含协议全部 {len(groups)} 类设备:")
    for group, pts in groups.items():
        print(f"    {group}{args.device_suffix}: {len(pts)} 点")
    print(f"  点位合计: {len(points)} (源码标识符, 不翻译)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
